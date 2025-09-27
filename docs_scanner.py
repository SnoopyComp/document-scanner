#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
docs_scanner.py
Scan multiple .docx files for one or more keywords and output match locations
with +/-N characters of context.

Usage examples:
  python docs_scanner.py -p "C:/docs" -k 딥러닝 -k 블록체인 -n 50 -i -o results.csv
  python docs_scanner.py -f "C:/docs/file1.docx" "C:/docs/file2.docx" -k 검색어

Notes:
- Supports .docx (Office Open XML). For legacy .doc, convert to .docx first.
- Matching is substring-based by default; enable -i for case-insensitive search.
"""

import argparse
import csv
import os
import re
import sys
from typing import Iterable, List, Tuple

try:
    from docx import Document  # python-docx
except ImportError:
    print("Missing dependency: python-docx\nInstall with: pip install python-docx", file=sys.stderr)
    sys.exit(1)


def iter_docx_text(docx_path: str) -> str:
    """Extract plain text from a .docx by joining paragraphs with newlines."""
    try:
        doc = Document(docx_path)
    except Exception as e:
        raise RuntimeError(f"Failed to open {docx_path}: {e}")
    paras = []
    for p in doc.paragraphs:
        # Include all runs by using paragraph.text. For tables, handle separately.
        paras.append(p.text)
    # Handle tables (optional but useful)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text for c in row.cells]
            if any(cells):
                paras.append("\t".join(cells))
    return "\n".join(paras)


def find_all_overlapping(pattern: re.Pattern, text: str) -> Iterable[re.Match]:
    """Yield overlapping matches."""
    i = 0
    while True:
        m = pattern.search(text, i)
        if not m:
            break
        yield m
        # Advance by 1 to allow overlaps
        i = m.start() + 1


def search_text(text: str, keywords: List[str], ignore_case: bool) -> List[Tuple[str, int, int]]:
    """
    Return list of (keyword, start, end) for all matches in text.
    """
    flags = re.IGNORECASE if ignore_case else 0
    results = []
    for kw in keywords:
        # Escape keyword to treat as literal substring
        pat = re.compile(re.escape(kw), flags)
        for m in find_all_overlapping(pat, text):
            results.append((kw, m.start(), m.end()))
    # Sort by position
    results.sort(key=lambda x: x[1])
    return results


def make_context(text: str, start: int, end: int, n: int) -> Tuple[str, str, str]:
    a = max(0, start - n)
    b = min(len(text), end + n)
    before = text[a:start]
    match = text[start:end]
    after = text[end:b]
    return before, match, after


def walk_docx_files(paths: List[str]) -> Iterable[str]:
    """
    Yield .docx files from provided paths. Each path can be a file or directory.
    """
    for p in paths:
        if os.path.isfile(p) and p.lower().endswith(".docx"):
            yield p
        elif os.path.isdir(p):
            for root, _, files in os.walk(p):
                for name in files:
                    if name.lower().endswith(".docx"):
                        yield os.path.join(root, name)
        else:
            # Skip non-existing or unsupported files silently
            continue


def human_pos(idx: int) -> int:
    """Convert 0-based index to 1-based for human-friendly reporting."""
    return idx + 1


def main():
    ap = argparse.ArgumentParser(description="Scan .docx files for keywords and output contexts.")
    grp = ap.add_mutually_exclusive_group(required=True)
    grp.add_argument("-p", "--path", help="Directory to scan recursively for .docx files")
    grp.add_argument("-f", "--files", nargs="+", help="One or more .docx files")

    ap.add_argument("-k", "--keyword", dest="keywords", action="append", required=True,
                    help="Keyword to search (can be repeated)")
    ap.add_argument("-n", "--context", type=int, default=50, help="Number of context characters on each side")
    ap.add_argument("-i", "--ignore-case", action="store_true", help="Case-insensitive search")
    ap.add_argument("-o", "--output", default="docx_keyword_results.csv", help="Output CSV path")

    args = ap.parse_args()

    paths = []
    if args.path:
        paths = [args.path]
    else:
        paths = args.files

    files = list(walk_docx_files(paths))
    if not files:
        print("No .docx files found in the given path(s).", file=sys.stderr)
        sys.exit(2)

    total_matches = 0
    # Prepare CSV
    with open(args.output, "w", newline="", encoding="utf-8-sig") as fp:
        w = csv.writer(fp)
        w.writerow(["file", "match_no", "pos_start_1based", "pos_end_1based", "keyword", "before", "match", "after", "excerpt"])

        for fpath in files:
            try:
                text = iter_docx_text(fpath)
            except Exception as e:
                print(e, file=sys.stderr)
                continue

            matches = search_text(text, args.keywords, args.ignore_case)
            for idx, (kw, s, e) in enumerate(matches, start=1):
                before, match, after = make_context(text, s, e, args.context)
                excerpt = f"...{before[-args.context:]}{match}{after[:args.context]}..."
                w.writerow([fpath, idx, human_pos(s), human_pos(e), kw, before, match, after, excerpt])
                total_matches += 1

            # Console summary per file
            print(f"{os.path.basename(fpath)}: {len(matches)} match(es)")

    print(f"\nDone. Total matches: {total_matches}. CSV saved to: {args.output}")


if __name__ == "__main__":
    main()
